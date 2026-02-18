#!/usr/bin/env python3
"""BMS CSVから電圧・電流プロット付きExcelを生成する。"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference


def _sort_vcell_columns(columns: Sequence[str]) -> List[str]:
    pattern = re.compile(r"^vcell(\d+)_v$")

    def key(name: str) -> int:
        match = pattern.match(name)
        return int(match.group(1)) if match else 9999

    return sorted(columns, key=key)


def _to_int_or_none(value: str) -> Optional[int]:
    value = value.strip()
    if value == "":
        return None
    return int(value)


def _to_float_or_none(value: str) -> Optional[float]:
    value = value.strip()
    if value == "":
        return None
    return float(value)


def _load_csv_rows(
    input_csv: Path,
    chain_id_filter: Optional[int],
    device_id_filter: Optional[int],
) -> tuple[List[str], List[Dict[str, str]]]:
    with input_csv.open("r", encoding="utf-8-sig", newline="") as fp:
        reader = csv.DictReader(fp)
        if reader.fieldnames is None:
            raise RuntimeError("CSV header is missing.")
        headers = list(reader.fieldnames)

        required = {"frame_index", "chain_id", "device_id", "current_a"}
        missing = [col for col in required if col not in headers]
        if missing:
            raise RuntimeError(
                f"Required columns are missing: {', '.join(missing)}"
            )

        rows: List[Dict[str, str]] = []
        for row in reader:
            try:
                chain_id = int(row["chain_id"])
                device_id = int(row["device_id"])
            except (TypeError, ValueError):
                continue

            if chain_id_filter is not None and chain_id != chain_id_filter:
                continue
            if device_id_filter is not None and device_id != device_id_filter:
                continue
            rows.append(row)

    return headers, rows


def _create_voltage_chart(
    ws_data,
    ws_charts,
    row_count: int,
    voltage_col_start: int,
    voltage_col_end: int,
) -> None:
    chart = LineChart()
    chart.title = "Cell Voltage Trend"
    chart.style = 2
    chart.y_axis.title = "Voltage [V]"
    chart.x_axis.title = "Frame Index"
    chart.height = 10
    chart.width = 18

    x_values = Reference(ws_data, min_col=1, min_row=2, max_row=row_count)
    for col in range(voltage_col_start, voltage_col_end + 1):
        y_values = Reference(ws_data, min_col=col, min_row=1, max_row=row_count)
        chart.add_data(y_values, titles_from_data=True)
    chart.set_categories(x_values)
    for series in chart.series:
        series.smooth = False
        series.marker.symbol = "circle"
        series.marker.size = 4

    ws_charts.add_chart(chart, "A1")


def _create_current_chart(
    ws_data,
    ws_charts,
    row_count: int,
    current_col: int,
) -> None:
    chart = LineChart()
    chart.title = "Pack Current Trend"
    chart.style = 2
    chart.y_axis.title = "Current [A]"
    chart.x_axis.title = "Frame Index"
    chart.height = 8
    chart.width = 18

    x_values = Reference(ws_data, min_col=1, min_row=2, max_row=row_count)
    y_values = Reference(ws_data, min_col=current_col, min_row=1, max_row=row_count)
    chart.add_data(y_values, titles_from_data=True)
    chart.set_categories(x_values)
    for series in chart.series:
        series.smooth = False
        series.marker.symbol = "circle"
        series.marker.size = 4

    ws_charts.add_chart(chart, "A24")


def build_workbook(
    headers: Sequence[str],
    rows: Sequence[Dict[str, str]],
    output_xlsx: Path,
) -> None:
    vcell_columns = _sort_vcell_columns([name for name in headers if name.startswith("vcell") and name.endswith("_v")])
    if not vcell_columns:
        raise RuntimeError("No voltage columns found (expected: vcellN_v).")

    data_columns = ["frame_index", "chain_id", "device_id", "current_a"] + vcell_columns

    wb = Workbook()
    ws_data = wb.active
    ws_data.title = "Data"
    ws_charts = wb.create_sheet("Charts")

    ws_data.append(data_columns)

    for row in rows:
        out_row: List[object] = []
        out_row.append(_to_int_or_none(row.get("frame_index", "")))
        out_row.append(_to_int_or_none(row.get("chain_id", "")))
        out_row.append(_to_int_or_none(row.get("device_id", "")))
        out_row.append(_to_float_or_none(row.get("current_a", "")))
        for col in vcell_columns:
            out_row.append(_to_float_or_none(row.get(col, "")))
        ws_data.append(out_row)

    ws_data.freeze_panes = "A2"

    row_count = len(rows) + 1
    current_col = data_columns.index("current_a") + 1
    voltage_col_start = data_columns.index(vcell_columns[0]) + 1
    voltage_col_end = data_columns.index(vcell_columns[-1]) + 1

    _create_voltage_chart(
        ws_data=ws_data,
        ws_charts=ws_charts,
        row_count=row_count,
        voltage_col_start=voltage_col_start,
        voltage_col_end=voltage_col_end,
    )
    _create_current_chart(
        ws_data=ws_data,
        ws_charts=ws_charts,
        row_count=row_count,
        current_col=current_col,
    )

    ws_charts["A45"] = "Source CSV columns are copied to 'Data'."
    ws_charts["A46"] = "Charts are generated for voltage (all cells) and current."

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate Excel template with voltage/current plots from BMS CSV."
    )
    parser.add_argument(
        "--input-csv",
        type=Path,
        required=True,
        help="Input CSV generated by bmschain_gui_serial_to_csv.py",
    )
    parser.add_argument(
        "--output-xlsx",
        type=Path,
        required=True,
        help="Output Excel file (.xlsx)",
    )
    parser.add_argument(
        "--chain-id",
        type=int,
        default=None,
        help="Optional filter for chain_id",
    )
    parser.add_argument(
        "--device-id",
        type=int,
        default=None,
        help="Optional filter for device_id",
    )
    return parser


def main() -> int:
    parser = build_argument_parser()
    args = parser.parse_args()

    headers, rows = _load_csv_rows(
        input_csv=args.input_csv,
        chain_id_filter=args.chain_id,
        device_id_filter=args.device_id,
    )
    if not rows:
        raise RuntimeError("No rows matched the input/filter conditions.")

    build_workbook(headers=headers, rows=rows, output_xlsx=args.output_xlsx)
    print(
        "[INFO] Wrote Excel template: "
        f"{args.output_xlsx} (rows={len(rows)}, chain_filter={args.chain_id}, device_filter={args.device_id})"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
